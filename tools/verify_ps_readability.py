"""Verify PS blank: contrast + print setup."""
from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"C:\Users\Gianmarco\corsi")
BLANK = ROOT / "assets/templates/ps_weekly_blank.xlsx"
SAMPLE = ROOT / "build/ps_fidelity/ooxml_sample.xlsx"


def main() -> None:
    wb = load_workbook(BLANK)
    ws = wb.active
    print("orient", ws.page_setup.orientation)
    print("fitW", ws.page_setup.fitToWidth, "fitH", ws.page_setup.fitToHeight)
    psp = ws.sheet_properties.pageSetUpPr
    print("fitToPage", psp.fitToPage if psp else None)
    print("margins", ws.page_margins.left, ws.page_margins.right)
    print("print_area", ws.print_area)
    print("sha16", hashlib.sha256(BLANK.read_bytes()).hexdigest()[:16])

    with zipfile.ZipFile(BLANK) as z:
        styles = z.read("xl/styles.xml").decode()
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        wbxml = z.read("xl/workbook.xml").decode()

    print("pageSetup", re.search(r"<pageSetup[^/]*/>", sheet).group(0))
    print("printArea", re.search(r"Print_Area.*?</definedName>", wbxml).group(0))

    fonts = re.findall(r"<font>.*?</font>", styles, flags=re.DOTALL)
    block = re.search(
        r"<!-- corsi-module-xfs -->(.*?)</cellXfs>", styles, flags=re.DOTALL
    ).group(1)
    mod_xfs = re.findall(r"<xf\b[^>]*>.*?</xf>", block, flags=re.DOTALL)
    labels = [
        (5, "mod6 lime"),
        (6, "mod7 amber"),
        (2, "mod3 cyan"),
        (12, "mod15 sky"),
        (0, "mod1 indigo"),
    ]
    for idx, label in labels:
        fid = int(re.search(r'fontId="(\d+)"', mod_xfs[idx]).group(1))
        color = re.search(r'color rgb="([^"]+)"', fonts[fid]).group(1)
        print(label, "font", color, "align left" if "left" in mod_xfs[idx] else "align?")

    if SAMPLE.exists():
        ws2 = load_workbook(SAMPLE).active
        for addr in ["D11", "D12"]:
            c = ws2[addr]
            fc = getattr(c.font.color, "rgb", None) if c.font.color else None
            fill = getattr(c.fill.fgColor, "rgb", None) if c.fill.fgColor else None
            print(
                "sample",
                addr,
                "val",
                c.value,
                "font",
                fc,
                "fill",
                fill,
                "size",
                c.font.size,
                "align",
                c.alignment.horizontal,
            )


if __name__ == "__main__":
    main()
