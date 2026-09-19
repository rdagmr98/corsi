"""Build assets/templates/ps_weekly_blank.xlsx from 66_PS EI sheet via openpyxl.

Then run: python tools/inject_module_styles.py
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill

SRC = Path(r"F:\66_PS 3° BTC B1 2025 EI+ CC - 2026.09.07.xlsx")
SRC_SHEET = "76S_09giugno_13giugno_EI"
OUT = Path(r"C:\Users\Gianmarco\corsi\assets\templates\ps_weekly_blank.xlsx")

DAY_BLOCKS = [(10, 17), (18, 25), (26, 33), (34, 41), (42, 44)]
LUNCH_OFFSET = 5
NO_FILL = PatternFill(fill_type=None)


def lesson_rows() -> list[int]:
    rows = []
    for start, end in DAY_BLOCKS:
        lunch = start + LUNCH_OFFSET if (end - start) == 7 else None
        for r in range(start, end + 1):
            if r != lunch:
                rows.append(r)
    return rows


def clear_cell_keep_border(cell) -> None:
    if isinstance(cell, MergedCell):
        return
    cell.value = None
    try:
        if cell.fill and cell.fill.patternType == "solid":
            rgb = None
            try:
                rgb = cell.fill.fgColor.rgb
            except Exception:
                pass
            if rgb and str(rgb).upper() in ("FFFFFF00", "FFFF00"):
                return
            cell.fill = NO_FILL
    except Exception:
        pass


def main() -> None:
    wb = load_workbook(SRC)
    for name in list(wb.sheetnames):
        if name != SRC_SHEET:
            del wb[name]
    ws = wb[SRC_SHEET]
    ws.title = "Settimana"

    for start, _ in DAY_BLOCKS:
        cell = ws.cell(start, 2)
        if not isinstance(cell, MergedCell):
            cell.value = None
    for r in lesson_rows():
        for c in range(4, 15):
            clear_cell_keep_border(ws.cell(r, c))

    ws["B5"].value = "{{COURSE_TITLE}}"
    for addr in ("D6", "J6", "D7"):
        cell = ws[addr]
        if not isinstance(cell, MergedCell):
            cell.value = None
    ws["B46"].value = "Direttore del corso: "

    for r in range(50, 66):
        for c in (2, 3, 9, 10):
            cell = ws.cell(r, c)
            if isinstance(cell, MergedCell):
                continue
            if cell.value is not None:
                cell.value = None

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("saved", OUT, OUT.stat().st_size)


if __name__ == "__main__":
    main()
