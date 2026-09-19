"""Structural comparison after OOXML rebuild — write report."""
from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(r"C:\Users\Gianmarco\corsi")
SRC = Path(r"F:\66_PS 3° BTC B1 2025 EI+ CC - 2026.09.07.xlsx")
BLANK = ROOT / "assets/templates/ps_weekly_blank.xlsx"
SAMPLE = ROOT / "build/ps_fidelity/ooxml_sample.xlsx"
OUT = ROOT / "build/ps_fidelity/comparison_report.md"


def merges(ws):
    return sorted(str(m) for m in ws.merged_cells.ranges)


def col_widths(ws, n=15):
    return {
        get_column_letter(i): ws.column_dimensions[get_column_letter(i)].width
        for i in range(1, n + 1)
    }


def zip_stats(path: Path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        styles = z.read("xl/styles.xml")
        sheets = [n for n in names if n.startswith("xl/worksheets/sheet")]
        sheet = z.read(sheets[0])
    return {
        "size": path.stat().st_size,
        "sha16": hashlib.sha256(path.read_bytes()).hexdigest()[:16],
        "entries": len(names),
        "styles_len": len(styles),
        "sheet_len": len(sheet),
        "ffff00": b"FFFF00" in styles.upper(),
        "has_shared": "xl/sharedStrings.xml" in names,
        "module_fill": b"FF6366F1" in styles.upper(),
    }


def header_probe(ws, label):
    lines = [f"### {label}"]
    for addr in ["B2", "B5", "B6", "D6", "B8", "C15", "D11", "G11", "I11"]:
        cell = ws[addr]
        try:
            fill = cell.fill.fgColor.rgb
        except Exception:
            fill = getattr(cell.fill.fgColor, "theme", None)
        lines.append(
            f"- {addr}: val={cell.value!r} bold={cell.font.bold} "
            f"size={cell.font.size} fill={fill}"
        )
    return "\n".join(lines)


def main():
    ei = load_workbook(SRC)["76S_09giugno_13giugno_EI"]
    blank = load_workbook(BLANK).active
    sample = load_workbook(SAMPLE).active
    sm, bm, om = merges(ei), merges(blank), merges(sample)
    ok = (
        sm == bm == om
        and col_widths(blank) == col_widths(ei) == col_widths(sample)
        and bool(sample["B6"].font.bold)
        and sample["B6"].font.size == 8
        and sample["C15"].fill.fgColor.rgb in ("FFFFFF00", "FFFF00")
        and "Modulo 15" in str(sample["D11"].value)
        and str(sample["D11"].fill.fgColor.rgb).upper().endswith("0EA5E9")  # module 15 palette
        and zip_stats(BLANK)["has_shared"]
        and zip_stats(BLANK)["module_fill"]
        and zip_stats(BLANK)["styles_len"] > 470000
    )
    lines = [
        "# PS Excel fidelity comparison",
        "",
        f"- Source EI sheet of `{SRC.name}`",
        f"- Blank `{BLANK}`",
        f"- Sample `{SAMPLE}`",
        "",
        "## Zip / styles",
        f"- SRC {zip_stats(SRC)}",
        f"- BLANK {zip_stats(BLANK)}",
        f"- SAMPLE {zip_stats(SAMPLE)}",
        "",
        f"## Merges EI={len(sm)} blank==EI={sm == bm} sample==EI={sm == om}",
        f"## Widths blank==EI={col_widths(blank) == col_widths(ei)} sample==EI={col_widths(sample) == col_widths(ei)}",
        "",
        header_probe(ei, "EI"),
        "",
        header_probe(blank, "BLANK"),
        "",
        header_probe(sample, "SAMPLE"),
        "",
        f"## Verdict: {'PASS' if ok else 'FAIL'}",
    ]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(OUT.read_text(encoding="utf-8"))


if __name__ == "__main__":
    main()
